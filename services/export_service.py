from pathlib import Path
from docx import Document

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = BASE_DIR / "outputs"
OUTPUT_DIR.mkdir(exist_ok=True)


def export_document(html_content: str, title: str, output_format: str, image_paths: list[str]) -> Path:
    safe_title = title.strip() or "document"
    file_name = safe_title.lower().replace(" ", "_")

    if output_format == "docx":
        output_path = OUTPUT_DIR / f"{file_name}.docx"
        doc = Document()
        doc.add_heading(safe_title, 0)
        for line in html_content.splitlines():
            if line.strip():
                doc.add_paragraph(line)
        doc.save(str(output_path))
        return output_path

    output_path = OUTPUT_DIR / f"{file_name}.html"
    output_path.write_text(html_content, encoding="utf-8")
    return output_path
