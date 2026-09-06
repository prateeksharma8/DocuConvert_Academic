#!/usr/bin/env python3
"""
Article Markdown to DOCX er
Run:
    pip install python-docx pillow
    python3 paper_.py
"""

import re, os, sys
from datetime import date
from pathlib import Path
from docx import Document
from docx.shared import Pt, Twips, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    from PIL import Image, ImageDraw
except Exception:
    Image = None
    ImageDraw = None

# paths - set defaults
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SRC_MD = os.path.join(BASE_DIR, "Q2ChangesFinal.md")
OUT_DOC = os.path.join(BASE_DIR, "Q2ChangesFinal.docx")
IMG_DIR   = os.path.join(BASE_DIR, "visualizations")
DOCX_NAME = "Q2ChangesFinal.docx"
TEMPLATE_DOC = os.path.join(BASE_DIR, "template_ES.docx")

PAGE_W = 12240; PAGE_H = 15840; MARGIN = 1440
CONTENT_W = PAGE_W - 2*MARGIN
COLUMN_GAP = 720
COLUMN_W = (CONTENT_W - COLUMN_GAP) // 2
BODY_W = COLUMN_W
BODY_W_EMU = int((BODY_W / 1440) * 914400)

# Match the supplied template's page geometry and column spacing
TEMPLATE_PAGE_W = 7772400
TEMPLATE_PAGE_H = 10058400
TEMPLATE_MARGIN = 914400
TEMPLATE_COLUMN_GAP = 720
TEMPLATE_CONTENT_W = TEMPLATE_PAGE_W - 2 * TEMPLATE_MARGIN
TEMPLATE_COLUMN_GAP_EMU = TEMPLATE_COLUMN_GAP * 635
TEMPLATE_BODY_W_EMU = (TEMPLATE_CONTENT_W - TEMPLATE_COLUMN_GAP_EMU) // 2
TABLE_FONT_SIZE = 8.5
TABLE_HEADER_FONT_SIZE = 8.5

def _tcPr(cell): return cell._tc.get_or_add_tcPr()
def set_cell_bg(cell, hex_color):
    shd=OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:color"),"auto"); shd.set(qn("w:fill"),hex_color); _tcPr(cell).append(shd)
def set_cell_borders(cell, hex_color="CCCCCC", sz=1):
    tcB=OxmlElement("w:tcBorders")
    for side in ("top","left","bottom","right"):
        el=OxmlElement(f"w:{side}"); el.set(qn("w:val"),"single"); el.set(qn("w:sz"),str(sz*8)); el.set(qn("w:space"),"0"); el.set(qn("w:color"),hex_color); tcB.append(el)
    _tcPr(cell).append(tcB)
def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcM=OxmlElement("w:tcMar")
    for side,val in (("top",top),("left",left),("bottom",bottom),("right",right)):
        el=OxmlElement(f"w:{side}"); el.set(qn("w:w"),str(val)); el.set(qn("w:type"),"dxa"); tcM.append(el)
    _tcPr(cell).append(tcM)
def compact_cell_margins(cell):
    set_cell_margins(cell, top=35, bottom=35, left=55, right=55)
def set_cell_width(cell, w):
    tcW=OxmlElement("w:tcW"); tcW.set(qn("w:w"),str(w)); tcW.set(qn("w:type"),"dxa"); _tcPr(cell).append(tcW)
def set_table_width(table, w):
    tbl=table._tbl; tblPr=tbl.find(qn("w:tblPr"))
    if tblPr is None: tblPr=OxmlElement("w:tblPr"); tbl.insert(0,tblPr)
    tblW=OxmlElement("w:tblW"); tblW.set(qn("w:w"),str(w)); tblW.set(qn("w:type"),"dxa"); tblPr.append(tblW)
def set_two_columns(section):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), str(TEMPLATE_COLUMN_GAP))
def set_one_column(section):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), "1")
def compact_table(table):
    table.autofit = False
    tbl=table._tbl; tblPr=tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr=OxmlElement("w:tblPr"); tbl.insert(0,tblPr)
    layout=OxmlElement("w:tblLayout"); layout.set(qn("w:type"),"fixed"); tblPr.append(layout)
def keep_table_rows_together(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
def page_number_field(run):
    r=run._r
    fc1=OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"),"begin"); r.append(fc1)
    it=OxmlElement("w:instrText"); it.text=" PAGE "; it.set(qn("xml:space"),"preserve"); r.append(it)
    fc2=OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"),"end"); r.append(fc2)
def toc_field(doc):
    """Insert TOC field after the last current paragraph — NOT append() which goes to end of body."""
    p=OxmlElement("w:p"); r=OxmlElement("w:r")
    fc1=OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"),"begin"); fc1.set(qn("w:dirty"),"true")
    it=OxmlElement("w:instrText"); it.set(qn("xml:space"),"preserve"); it.text=' TOC \\o "1-3" \\h \\z \\u '
    fc2=OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"),"end")
    r.append(fc1); r.append(it); r.append(fc2); p.append(r)
    body=doc.element.body; last_p=body.findall(qn("w:p"))[-1] if body.findall(qn("w:p")) else None
    if last_p is not None: last_p.addnext(p)
    else: body.insert(0, p)
def para_top_border(para, hex_color, sz_pt, space=4):
    pPr=para._p.get_or_add_pPr()
    pBdr=pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr=OxmlElement("w:pBdr"); pPr.append(pBdr)
    top=OxmlElement("w:top")
    top.set(qn("w:val"),"single"); top.set(qn("w:sz"),str(sz_pt*8)); top.set(qn("w:space"),str(space)); top.set(qn("w:color"),hex_color)
    pBdr.append(top)

def para_bottom_border(para, hex_color, sz_pt, space=4):
    pPr=para._p.get_or_add_pPr()
    pBdr=pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr=OxmlElement("w:pBdr"); pPr.append(pBdr)
    bot=OxmlElement("w:bottom")
    bot.set(qn("w:val"),"single"); bot.set(qn("w:sz"),str(sz_pt*8)); bot.set(qn("w:space"),str(space)); bot.set(qn("w:color"),hex_color)
    pBdr.append(bot)
def spacing(para, before_pt=0, after_pt=8):
    pPr=para._p.get_or_add_pPr(); spc=OxmlElement("w:spacing")
    spc.set(qn("w:before"),str(int(before_pt*20))); spc.set(qn("w:after"),str(int(after_pt*20))); pPr.append(spc)
def tight_table_spacing(para, line_dxa=185):
    pPr=para._p.get_or_add_pPr(); spc=OxmlElement("w:spacing")
    spc.set(qn("w:before"),"0"); spc.set(qn("w:after"),"0")
    spc.set(qn("w:line"),str(line_dxa)); spc.set(qn("w:lineRule"),"exact")
    pPr.append(spc)
def indent_para(para, left_dxa):
    pPr=para._p.get_or_add_pPr(); ind=OxmlElement("w:ind"); ind.set(qn("w:left"),str(left_dxa)); pPr.append(ind)
def rfmt(run, font="Times New Roman", sz=10, bold=False, italic=False, color=None):
    run.font.name=font; run.font.size=Pt(sz); run.font.bold=bold; run.font.italic=italic
    if color: run.font.color.rgb=color

def clear_document_body(doc):
    body_el = doc.element.body
    for child in list(body_el):
        if child.tag != qn("w:sectPr"):
            body_el.remove(child)

def safe_set_style(paragraph, style_name):
    """Try to set paragraph style, fall back to 'Normal' if style doesn't exist"""
    try:
        paragraph.style = style_name
    except KeyError:
        try:
            paragraph.style = 'Normal'
        except KeyError:
            pass  # Use default

def ensure_document_styles(doc):
    """Preserve the template's built-in styles and only rely on them."""
    return None


def _section_indices(doc):
    body = doc.element.body
    return [idx for idx, child in enumerate(list(body)) if child.tag == qn("w:sectPr")]


def clear_section_content(doc, section_index):
    body = doc.element.body
    children = list(body)
    positions = _section_indices(doc)
    if not positions or section_index < 0 or section_index >= len(positions):
        return
    start = 0 if section_index == 0 else positions[section_index - 1] + 1
    end = positions[section_index]
    for idx in range(end - 1, start - 1, -1):
        body.remove(children[idx])


def insert_element_into_section(doc, section_index, element):
    body = doc.element.body
    children = list(body)
    positions = _section_indices(doc)
    if not positions or section_index < 0 or section_index >= len(positions):
        return
    if element in children:
        body.remove(element)
    insert_at = positions[section_index]
    body.insert(insert_at, element)


def create_paragraph(doc, style="Body Text", section_index=None):
    p = doc.add_paragraph(style=style)
    if section_index is not None:
        insert_element_into_section(doc, section_index, p._p)
    return p


def find_template_document():
    """Use the supplied Word template for styles and page layout while keeping the body content fresh."""
    candidates = [
        os.path.join(BASE_DIR, "template_ES.docx"),
        os.path.join(BASE_DIR, "template.docx"),
        os.path.join(os.path.dirname(__file__), "templates", "template_ES.docx"),
        os.path.expanduser("~/Desktop/template_ES.docx"),
    ]
    for candidate in candidates:
        if os.path.exists(candidate):
            return candidate
    return None


def setup_doc():
    template_path = find_template_document()

    if template_path:
        doc = Document(template_path)
        remove_template_placeholder_paragraphs(doc)
    else:
        doc = Document()

    # Keep first section ONE COLUMN
    first = doc.sections[0]

    first.page_width = Emu(TEMPLATE_PAGE_W)
    first.page_height = Emu(TEMPLATE_PAGE_H)
    first.left_margin = Emu(TEMPLATE_MARGIN)
    first.right_margin = Emu(TEMPLATE_MARGIN)
    first.top_margin = Emu(TEMPLATE_MARGIN)
    first.bottom_margin = Emu(TEMPLATE_MARGIN)

    set_one_column(first)

    return doc

_STRIP_NUM=re.compile(r"^[\d]+(\.\d+)*\.?\s+")
def clean_h(text): return _STRIP_NUM.sub("",text.strip()).strip("#").strip()
def heading_text(text): return text.strip().strip("#").strip()

def sanitize_text(text):
    if text is None:
        return ""
    text = str(text)
    text = text.replace("\x00", "")
    text = ''.join(ch for ch in text if ch.isprintable() or ch in "\n\t")
    return text


SECTION_COUNTER = 0
SUBSECTION_COUNTER = 0

def add_h1(doc, text, numbered=True, all_caps=True):
    global SECTION_COUNTER, SUBSECTION_COUNTER
    SECTION_COUNTER += 1
    SUBSECTION_COUNTER = 0
    roman = ["I","II","III","IV","V","VI","VII","VIII","IX","X"]
    heading = clean_h(sanitize_text(text))
    if numbered and SECTION_COUNTER <= len(roman):
        heading = f"{roman[SECTION_COUNTER-1]}. {heading}"
    para = doc.add_paragraph(style="Heading 1")
    run = para.add_run(heading.upper() if all_caps else heading)
    rfmt(run, sz=10, bold=True)
    return para

def add_h2(doc, text):
    global SUBSECTION_COUNTER
    SUBSECTION_COUNTER += 1
    letters = ["A","B","C","D","E","F","G","H","I","J"]
    prefix = f"{letters[SUBSECTION_COUNTER-1]}. " if SUBSECTION_COUNTER <= len(letters) else ""
    para = doc.add_paragraph(style="Heading 2")
    run = para.add_run(prefix + heading_text(text).upper())
    rfmt(run, sz=10, bold=True)
    return para

def add_h3(doc, text):
    para = doc.add_paragraph(style="Heading 3")
    run = para.add_run(heading_text(text).upper())
    rfmt(run, sz=10, bold=True)
    return para

def add_h4(doc, text):
    para = doc.add_paragraph(style="Heading 4")
    run = para.add_run(heading_text(text).upper())
    rfmt(run, sz=10, bold=True)
    return para

_INLINE=re.compile(r'\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`')
def add_inline(para,text):
    text = sanitize_text(text)
    pos=0
    for m in _INLINE.finditer(text):
        if m.start()>pos: para.add_run(text[pos:m.start()])
        if m.group(1): para.add_run(m.group(1)).bold=True
        elif m.group(2): para.add_run(m.group(2)).italic=True
        elif m.group(3): para.add_run(m.group(3)).font.name="Courier New"
        pos=m.end()
    if pos<len(text): para.add_run(text[pos:])

def body(doc, text, indent_dxa=0):

    p = add_style_safe_paragraph(doc, "Body Text", "Normal")

    if indent_dxa:
        indent_para(p, indent_dxa)

    add_inline(p, text)

    return p

def compact_reference_body(doc,text,indent_dxa=0):
    return add_reference(doc, text, indent_dxa)

def add_reference(doc, text, indent_dxa=0):
    """Add a bibliography entry using the template's reference style."""
    p = add_style_safe_paragraph(doc, "references", "Normal")

    if indent_dxa:
        indent_para(p, indent_dxa)

    # remove duplicated numbering like "1. "
    text = re.sub(r'^\d+\.\s*', '', text)

    add_inline(p, text)

    return p

def list_para(doc,text):
    p=doc.add_paragraph(style="bullet list")
    add_inline(p,text); return p

def fresh_decimal_num_id(doc):
    """Create a decimal list instance that restarts at 1 for each list group."""
    numbering = doc.part.numbering_part._element
    existing = [int(num.get(qn("w:numId"))) for num in numbering.xpath("./w:num")]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_id = OxmlElement("w:abstractNumId")
    abstract_id.set(qn("w:val"), "3")  # template_ES decimal format: %1.
    num.append(abstract_id)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id

def numbered_para(doc, text, num_id):
    """Add a decimal list item using the template's decimal numbering definition."""
    p = doc.add_paragraph(style="Normal")
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    identifier = OxmlElement("w:numId")
    identifier.set(qn("w:val"), str(num_id))
    numpr.extend((level, identifier))
    ppr.append(numpr)
    add_inline(p, text)
    return p

BOX=set("│┌┐└┘─═╔╗╚╝▼╠╣╦╩╬")
def is_diagram(lines): return any(any(c in BOX for c in ln) for ln in lines)
def add_code_block(doc,lines):
    for ln in lines:
        p=doc.add_paragraph(style="Body Text")
        p.add_run(ln if ln else " ").font.name="Courier New"

def parse_table(lines):
    rows=[]
    for ln in lines:
        if re.match(r"^\s*\|?[-:| ]+\|?\s*$",ln): continue
        rows.append([c.strip() for c in ln.strip().strip("|").split("|")])
    return rows

def _cell_write(cell, raw_txt, sz=TABLE_FONT_SIZE, bold=False):
    """Split <br> tags into separate Word paragraphs within a cell."""
    raw_txt=re.sub(r'(<br>\s*)+','<br>',raw_txt); segments=[s.strip() for s in raw_txt.split('<br>')]
    for ep in list(cell.paragraphs): ep._p.getparent().remove(ep._p)
    for seg in segments:
        cp=OxmlElement("w:p"); pPr=OxmlElement("w:pPr")
        spc=OxmlElement("w:spacing"); spc.set(qn("w:before"),"0"); spc.set(qn("w:after"),"0"); pPr.append(spc); cp.append(pPr)
        jc=OxmlElement("w:jc"); jc.set(qn("w:val"),"left"); pPr.append(jc)
        cr=OxmlElement("w:r"); cPr=OxmlElement("w:rPr")
        fn=OxmlElement("w:rFonts"); fn.set(qn("w:ascii"),"Arial"); fn.set(qn("w:hAnsi"),"Arial")
        sz_el=OxmlElement("w:sz"); sz_el.set(qn("w:val"),str(sz*2)); szCs=OxmlElement("w:szCs"); szCs.set(qn("w:val"),str(sz*2))
        col=OxmlElement("w:color"); col.set(qn("w:val"),"333333")
        cPr.append(fn); cPr.append(sz_el); cPr.append(szCs); cPr.append(col)
        if bold: cPr.append(OxmlElement("w:b"))
        cr.append(cPr)
        ct=OxmlElement("w:t"); ct.set(qn("xml:space"),"preserve")
        ct.text=re.sub(r'\*\*(.+?)\*\*',r'\1',seg) if seg else " "
        cr.append(ct); cp.append(cr); cell._tc.append(cp)

def generic_table(doc, lines):
    rows=parse_table(lines)
    if not rows: return
    nc=max(len(r) for r in rows)
    tbl=doc.add_table(len(rows),nc, style="Table Grid")
    set_table_width(tbl,BODY_W); compact_table(tbl)
    keep_table_rows_together(tbl)
    for ri,row in enumerate(rows):
        for ci in range(nc):
            cell=tbl.rows[ri].cells[ci]
            paragraph=cell.paragraphs[0]
            paragraph.clear()
            safe_set_style(paragraph, "table head" if ri == 0 else "table copy")
            add_inline(paragraph, re.sub(r'\*\*(.+?)\*\*',r'\1',row[ci].strip() if ci<len(row) else ""))

def create_placeholder_image(img_path, label):
    if Image is None or ImageDraw is None:
        return None
    out_dir = os.path.dirname(img_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    width, height = 1200, 800
    img = Image.new("RGB", (width, height), (248, 248, 248))
    draw = ImageDraw.Draw(img)
    draw.rectangle((20, 20, width - 20, height - 20), outline=(180, 180, 180), width=6)
    draw.text((50, 50), (label or "Figure placeholder")[:90], fill=(70, 70, 70))
    img.save(img_path)
    return img_path


def add_figure(doc, img_path, width_emu, caption_text):
    if not img_path:
        body(doc, "[IMAGE NOT FOUND]"); return
    resolved = img_path
    if not os.path.isabs(resolved):
        resolved = resolve_image_path(resolved)
    if not resolved or not os.path.exists(resolved):
        placeholder_path = os.path.join(IMG_DIR, f"{os.path.splitext(os.path.basename(img_path))[0]}_placeholder.png")
        if not os.path.exists(placeholder_path):
            create_placeholder_image(placeholder_path, caption_text or os.path.basename(img_path))
        resolved = placeholder_path if os.path.exists(placeholder_path) else None
    if not resolved:
        body(doc, f"[IMAGE NOT FOUND: {img_path}]"); return
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(resolved, width=Emu(width_emu))

    caption_text = re.sub(
        r"^Figure\s+\d+\s*:\s*",
        "",
        caption_text,
        flags=re.I,
    )

    cap = add_style_safe_paragraph(doc, "figure caption", "Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    cap.add_run("Fig. ").bold = True
    cap.add_run(caption_text)


def resolve_image_path(path):
    """Resolve image path from markdown, including workspace-wide fallback search."""
    path = (path or "").strip()
    if not path:
        return None
    if os.path.isabs(path):
        return path
    normalized = path.replace("\\", "/")
    if normalized.startswith("visualizations/"):
        normalized = normalized[len("visualizations/"):]
    candidates = []
    src_dir = os.path.dirname(os.path.abspath(SRC_MD))
    candidates.extend([
        os.path.join(src_dir, normalized),
        os.path.join(BASE_DIR, normalized),
        os.path.join(BASE_DIR, "visualizations", normalized),
        os.path.join(BASE_DIR, "analysis_outputs", normalized),
    ])
    for match in Path(BASE_DIR).rglob(os.path.basename(normalized)):
        if match.is_file():
            candidates.append(str(match))
    for candidate in candidates:
        if os.path.exists(candidate):
            return candidate
    return None

def is_equation_line(text):
    stripped = text.strip()
    if not stripped or stripped.startswith(("#", "- ", "* ", ">", "|", "```")):
        return False
    has_math_term = bool(re.search(r"\b(beta|epsilon|Delta|Mortality|Incidence|GDP|LifeExp|Fertility|Urbanization|SRBImbalance|alpha|gamma|mu|sigma)\b", stripped, re.I))
    return "=" in stripped and has_math_term


def format_equation_text(text):
    eq = text.strip()
    eq = eq.replace("Delta ", "Δ ").replace("Delta", "Δ")
    eq = eq.replace("epsilon", "ε")
    eq = re.sub(r"\bbeta_(\d+)\b", lambda m: f"β{''.join('₀₁₂₃₄₅₆₇₈₉'[int(ch)] for ch in m.group(1))}", eq)
    return eq


def render_equation_block(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(format_equation_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.font.italic = True


FIGS = {
    "figure_1": {"file":"01_main_model_explanatory/01_main_model_explanatory_linear_regression_model_comparison.png","cx":5_943_600,
        "caption":"Figure 1: Nested model R2 comparison"},
    "figure_2": {"file":"01_main_model_explanatory/01_main_model_explanatory_linear_regression_standardized_coefficients.png","cx":5_943_600,
        "caption":"Figure 2: Standardized coefficient ranking"},
    "figure_3": {"file":"02_validation_leave_one_country_out/02_validation_leave_one_country_out_generalization_actual_vs_predicted.png","cx":5_943_600,
        "caption":"Figure 3: Leave-one-country-out actual vs predicted mortality"},
    "figure_4": {"file":"03_supplementary_forecasting/03_supplementary_catboost_xgboost_forecasting_2016_2023_permutation_importance.png","cx":5_943_600,
        "caption":"Figure 4: Supplementary forecasting permutation importance"},
}

def split_sections(md_text):
    """
    Line-number based splitter. If source MD is edited, update bounds.
    Find new line numbers with: grep -n "^# \\|^## \\|^### " source.md
    """
    lines=md_text.split("\n")
    bounds=[
        (0,"preamble"),(18,"admin"),(40,"field"),(50,"objects"),(70,"summary"),
        (82,"drawings"),(116,"background"),(150,"priorart"),(332,"detailed_main"),
        (618,"sec422"),(675,"additional"),(687,"commercial_body"),
        (698,"sec421"),(713,"sec423"),(745,"sec43"),(799,"claims"),
        (len(lines),None),
    ]
    sections={}
    for idx in range(len(bounds)-1):
        start_line,key=bounds[idx]; end_line,_=bounds[idx+1]
        if key: sections[key]=lines[start_line:end_line]
    return sections



def render_claims(doc, lines):
    i=0; n=len(lines)
    while i < n:
        ln=lines[i].strip()
        if not ln or re.match(r'^#+\s*CLAIMS',ln,re.I): i+=1; continue
        m_claim=re.match(r'^\*\*(\d+)\.\*\*\s*(.*)',ln)
        if m_claim:
            num=m_claim.group(1); rest=m_claim.group(2)
            preamble=" ".join(bl for bl in lines[i+1:i+3] if bl)
            p=doc.add_paragraph(); spacing(p,6,4)
            rfmt(p.add_run(f"{num}.  "),bold=True)
            if preamble: add_inline(p,preamble)
            i+=1; continue
        p=doc.add_paragraph(); spacing(p,0,4); add_inline(p,ln); i+=1


def render_lines(doc, lines, skip_first_h1=True, insert_figure_after=None):
    if insert_figure_after is None: insert_figure_after={}
    i=0; n=len(lines); first_h1_skipped=not skip_first_h1; in_references=False
    def maybe_insert_fig(text):
        for trigger,fig_key in list(insert_figure_after.items()):
            if trigger in text and fig_key in FIGS:
                fig=FIGS[fig_key]; add_figure(doc,os.path.join(IMG_DIR,fig["file"]),fig["cx"],fig["caption"])
                del insert_figure_after[trigger]; break
    while i<n:
        ln=lines[i]; stripped=ln.strip()
        if not stripped: i+=1; continue
        if re.match(r'^-{3,}$',stripped): i+=1; continue
        if is_equation_line(stripped):
            render_equation_block(doc, stripped)
            i+=1; continue
        img_match=re.match(r'^!\[(.*?)\]\((.*?)\)\s*$', stripped)
        if img_match:
            caption=img_match.group(1).strip() or "Figure"
            img_path=resolve_image_path(img_match.group(2).strip())
            add_figure(doc,img_path,BODY_W_EMU,caption)
            i+=1; continue
        if stripped.startswith("```"):
            fence=[]; i+=1
            while i<n and not lines[i].strip().startswith("```"):
                fence.append(lines[i].rstrip("\n"))
                i+=1
            i+=1
            add_code_block(doc,fence)
            continue
        if "|" in stripped and stripped.startswith("|"):
            tbl_lines=[]
            while i<n and "|" in lines[i] and lines[i].strip().startswith("|"): tbl_lines.append(lines[i]); i+=1
            generic_table(doc,tbl_lines); continue
        if stripped.startswith("#### "):
            in_references=False; add_h4(doc,stripped[5:]); i+=1; continue
        if stripped.startswith("### "):
            in_references=False; add_h3(doc,stripped[4:]); i+=1; continue
        if stripped.startswith("## "):
            heading=heading_text(stripped[3:])
            in_references=(heading.lower()=="references")
            add_h2(doc,stripped[3:]); i+=1; continue
        if stripped.startswith("# ") and not stripped.startswith("## "):
            heading=heading_text(stripped[2:])
            if not first_h1_skipped:
                first_h1_skipped=True
            in_references=(heading.lower()=="references")
            add_h1(doc,stripped[2:])
            i+=1; continue
        if stripped.startswith("> "):
            p=body(doc,re.sub(r'^> ?','',stripped),720); maybe_insert_fig(stripped); i+=1; continue
        if stripped.startswith("- ") or stripped.startswith("* "):
            while i<n:
                cur=lines[i]; cur_s=cur.strip()
                if not cur_s: break
                if cur_s.startswith("- ") or cur_s.startswith("* "):
                    m=re.match(r'^[-*]\s+(.*)',cur_s)
                    if m: list_para(doc,m.group(1))
                    i+=1; continue
                break
            continue
        if re.match(r'^\d+\.\s+', stripped):
            number_id = fresh_decimal_num_id(doc)
            while i<n:
                cur_s=lines[i].strip()
                m=re.match(r'^\d+\.\s+(.*)', cur_s)
                if not m: break
                if in_references:
                    compact_reference_body(doc,cur_s,indent_dxa=240)
                else:
                    numbered_para(doc,m.group(1),number_id)
                i+=1
            continue
        para_lines=[]
        while i<n and lines[i].strip() and not lines[i].strip().startswith("#") and not lines[i].strip().startswith("```") and not ("|" in lines[i] and lines[i].strip().startswith("|")) and not lines[i].strip().startswith("- ") and not lines[i].strip().startswith("* ") and not re.match(r'^-{3,}$',lines[i].strip()):
            para_lines.append(lines[i].strip()); i+=1
        if para_lines:
            full=" ".join(para_lines); body(doc,full); maybe_insert_fig(full)
        continue


def extract_title(md_text):
    for ln in md_text.splitlines():
        if ln.startswith("# "):
            return clean_h(ln[2:])
    return "Final Research Article"


def split_frontmatter(md_text):
    lines=md_text.splitlines()
    title=extract_title(md_text)
    abstract_lines=[]
    body_start=0
    in_abstract=False
    for idx, line in enumerate(lines):
        stripped=line.strip()
        if re.fullmatch(r"#{1,2}\s+Abstract", stripped, re.I):
            in_abstract=True
            continue
        if in_abstract and re.match(r"#{1,2}\s+", stripped):
            body_start=idx
            break
        if in_abstract and stripped:
            abstract_lines.append(stripped)
    if body_start == 0:
        for idx, line in enumerate(lines):
            if line.strip().startswith("## Introduction"):
                body_start=idx
                break
    return title, abstract_lines, lines[body_start:]


def remove_template_placeholder_paragraphs(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_style_safe_paragraph(doc, style_name, fallback="Normal"):
    try:
        return doc.add_paragraph(style=style_name)
    except KeyError:
        return doc.add_paragraph(style=fallback)


def add_template_frontmatter(doc, title, abstract_lines):

    title_para = add_style_safe_paragraph(doc, "Title")
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(title.upper())
    rfmt(run, sz=16, bold=True)

    author = add_style_safe_paragraph(doc, "Normal")
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = author.add_run("Author Name\nIndependent Researcher\nCity, Country\nemail@example.com")
    rfmt(r1, sz=10)

    # ---------- START TWO COLUMNS ----------
    new_section = doc.add_section(WD_SECTION.CONTINUOUS)
    new_section.page_width = Emu(TEMPLATE_PAGE_W)
    new_section.page_height = Emu(TEMPLATE_PAGE_H)
    new_section.left_margin = Emu(TEMPLATE_MARGIN)
    new_section.right_margin = Emu(TEMPLATE_MARGIN)
    new_section.top_margin = Emu(TEMPLATE_MARGIN)
    new_section.bottom_margin = Emu(TEMPLATE_MARGIN)
    set_two_columns(new_section)

    h = doc.add_paragraph()
    para_top_border(h, "000000", 1, 4)
    run_h = h.add_run("ABSTRACT")
    rfmt(run_h, sz=10, bold=True)

    for line in abstract_lines:
        if line.strip():
            p = doc.add_paragraph(style="Body Text")
            add_inline(p, line)
            for run in p.runs: rfmt(run, sz=10)

    kw = doc.add_paragraph(style="Body Text")
    r_kw = kw.add_run("KEYWORDS—")
    rfmt(r_kw, sz=10, bold=True)
    r_val = kw.add_run(" breast cancer, mortality, sex ratio at birth, machine learning")
    rfmt(r_val, sz=10)
    para_bottom_border(kw, "000000", 1, 4)


def study_manuscript_output(path):
    return path


def add_cover_letter(doc, title):
    today = date.today().strftime("%B %d, %Y").replace(" 0", " ")
    p=doc.add_paragraph(); spacing(p,0,12)
    rfmt(p.add_run("Cover Letter"),sz=16,bold=True)
    para_bottom_border(p,"2E75B6",2,4)
    for text in [today, "Author Name", "Email: email@example.com", "Dear Editor,", f'Please consider the manuscript titled "{title}" for review.']:
        body(doc,text)


def build(md_text):
    doc=setup_doc()
    title, abstract_lines, body_lines = split_frontmatter(md_text)
    add_template_frontmatter(doc,title,abstract_lines)
    render_lines(doc, body_lines, skip_first_h1=True)
    return doc


def build_cover_letter(md_text):
    global _doc_ref
    doc=setup_doc(); _doc_ref=doc
    set_one_column(doc.sections[0])
    title=extract_title(md_text)
    add_cover_letter(doc,title)
    return doc

if __name__ == "__main__":
    if len(sys.argv)>=2: SRC_MD=sys.argv[1]
    if len(sys.argv)>=3: OUT_DOC=study_manuscript_output(sys.argv[2])
    if len(sys.argv)>=4: IMG_DIR=sys.argv[3]
    OUT_DOC=study_manuscript_output(OUT_DOC)
    print(f"Source : {SRC_MD}"); print(f"Output : {OUT_DOC}"); print(f"Images : {IMG_DIR}"); print()
    if not os.path.exists(SRC_MD): print(f"ERROR: {SRC_MD}"); sys.exit(1)
    print("Reading..."); md_text=open(SRC_MD,encoding="utf-8").read()
    print("Building..."); doc=build(md_text)
    os.makedirs(os.path.dirname(os.path.abspath(OUT_DOC)),exist_ok=True)
    doc.save(OUT_DOC); size_kb=os.path.getsize(OUT_DOC)//1024
    print(f"Done  ({size_kb} KB)  ->  {OUT_DOC}")
