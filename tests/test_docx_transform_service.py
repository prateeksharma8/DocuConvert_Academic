import tempfile
import unittest
from pathlib import Path

from docx import Document

from services.docx_transform_service import transform_docx


class TestDocxTransformService(unittest.TestCase):
    def test_source_content_is_written_using_template_layout(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source_path = root / "source.docx"
            template_path = root / "template.docx"
            output_path = root / "formatted.docx"

            source = Document()
            source.add_heading("Source title", level=1)
            source.add_paragraph("Source body")
            source.save(source_path)

            template = Document()
            template.sections[0].left_margin = 123456
            template.add_paragraph("Template placeholder")
            template.save(template_path)

            transform_docx(source_path, template_path, output_path)
            result = Document(output_path)

            text = " ".join(paragraph.text for paragraph in result.paragraphs)
            self.assertIn("Source title", text)
            self.assertIn("Source body", text)
            self.assertNotIn("Template placeholder", text)
            self.assertAlmostEqual(result.sections[0].left_margin, 123456, delta=500)


if __name__ == "__main__":
    unittest.main()